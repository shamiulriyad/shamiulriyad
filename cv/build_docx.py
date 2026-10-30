# Builds an editable Word copy of the CV from the HTML source, so both stay in sync.
# Usage: python3 cv/build_docx.py [basename]   (needs python-docx and beautifulsoup4)
import sys
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

HERE = Path(__file__).parent
NAME = sys.argv[1] if len(sys.argv) > 1 else "Md_Shamiul_Islam_Riyad_Software_Engineer_Intern_CV"
LINK = RGBColor(0x1A, 0x4F, 0x8B)
MUTED = RGBColor(0x44, 0x44, 0x44)
FONT = "Calibri"
TEXT_WIDTH = Mm(210 - 2 * 14)


def text(el):
    return " ".join(el.get_text().split())


def spacing(p, before=0, after=0):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0


def add_hyperlink(p, url, label, size=None):
    rel = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1A4F8B")
    props.append(color)
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        props.append(sz)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = label
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    p._p.append(link)


def bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "444444")):
        bottom.set(qn(k), v)
    borders.append(bottom)
    pPr.append(borders)


def right_tab(p):
    p.paragraph_format.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)


def bullets(doc, ul):
    for li in ul.find_all("li"):
        p = doc.add_paragraph(style="List Bullet")
        spacing(p, after=1)
        p.paragraph_format.left_indent = Mm(4.5)
        label = li.find("strong")
        if label:
            p.add_run(text(label)).bold = True
            p.add_run(text(li)[len(text(label)):])
        else:
            p.add_run(text(li))


soup = BeautifulSoup((HERE / f"{NAME}.html").read_text(encoding="utf-8"), "html.parser")
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin, sec.bottom_margin = Mm(11), Mm(8)
sec.left_margin = sec.right_margin = Mm(14)
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(10)
doc.core_properties.author = text(soup.find("h1"))
doc.core_properties.title = soup.title.string

header = soup.find("header")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
spacing(p, after=2)
r = p.add_run(text(header.find("h1")))
r.bold, r.font.size = True, Pt(20)
tagline = header.find(class_="tagline")
if tagline:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    spacing(p, after=2)
    r = p.add_run(text(tagline))
    r.font.size, r.font.color.rgb = Pt(10.5), RGBColor(0x33, 0x33, 0x33)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
spacing(p, after=4)
for i, a in enumerate(header.select(".contact a")):
    if i:
        p.add_run("  |  ").font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    add_hyperlink(p, a["href"], text(a), size=9)

for section in soup.find_all("section"):
    p = doc.add_paragraph()
    spacing(p, before=6, after=3)
    r = p.add_run(text(section.find("h2")).upper())
    r.bold, r.font.size = True, Pt(10.5)
    bottom_border(p)

    if section.find(class_="skills"):
        for line in section.find_all("p", class_="skill"):
            label = text(line.find("strong"))
            p = doc.add_paragraph()
            spacing(p, after=1)
            p.paragraph_format.left_indent = Mm(3.5)
            p.paragraph_format.first_line_indent = Mm(-3.5)
            p.add_run(label).bold = True
            p.add_run(text(line)[len(label):])
        continue

    projects = section.find_all("div", class_="project")
    if projects:
        for i, proj in enumerate(projects):
            title = proj.find(class_="title")
            what = title.find(class_="what")
            p = doc.add_paragraph()
            spacing(p, before=4 if i else 0)
            right_tab(p)
            name = title.find(string=True, recursive=False).strip()
            r = p.add_run(name + " ")
            r.bold, r.font.size = True, Pt(10.5)
            p.add_run(text(what)).font.size = Pt(10.5)
            p.add_run("\t")
            a = proj.find("a", class_="link")
            add_hyperlink(p, a["href"], text(a), size=9.5)
            p = doc.add_paragraph()
            spacing(p, after=1)
            r = p.add_run(text(proj.find(class_="stack")))
            r.italic, r.font.size, r.font.color.rgb = True, Pt(9.5), MUTED
            bullets(doc, proj.find("ul"))
        continue

    row = section.find("div", class_="row")
    if row:
        left, right = row.find_all("div", recursive=False)
        p = doc.add_paragraph()
        spacing(p)
        right_tab(p)
        p.add_run(text(left.find("strong"))).bold = True
        p.add_run(text(left)[len(text(left.find("strong"))):])
        p.add_run("\t")
        p.add_run(text(right)).font.color.rgb = MUTED
        for note in section.find_all("p", class_="note"):
            p = doc.add_paragraph(text(note))
            spacing(p, before=1)
        continue

    if section.find("ul"):
        bullets(doc, section.find("ul"))
    else:
        p = doc.add_paragraph(text(section.find("p")))
        spacing(p)

out = HERE / f"{NAME}.docx"
doc.save(out)
print(f"Wrote cv/{out.name}")
