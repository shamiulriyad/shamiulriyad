# Builds the full assignment "Is GDP Growth an Adequate Measure of Economic Well-Being?" as a Word file.
# Usage:  python3 gdp-assignment/make_figures.py      (charts + cover image)
#         python3 gdp-assignment/build_assignment.py  (two passes: lays out, reads page numbers, fills the contents lists)
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

from content import BODY, ABSTRACT, ABBREVIATIONS, REFERENCES, GLOSSARY

HERE = Path(__file__).parent
FIG = HERE / "figures"
OUT = HERE / "GDP_Growth_and_Economic_Well-Being_Assignment.docx"

NAVY = RGBColor(0x1F, 0x38, 0x64)
NAVY_HEX = "1F3864"
PALE_HEX = "E8EDF6"
TEXT = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x59, 0x59, 0x59)
FONT = "Times New Roman"
TEXT_W = Mm(210 - 2 * 25)  # A4 with 25 mm side margins


# ---------------------------------------------------------------- low-level helpers
def set_font(run, size=None, bold=None, italic=None, color=None, font=FONT, caps=None):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    if caps is not None:
        run.font.all_caps = caps
    return run


def add_runs(p, text, size=None, color=None, italic=None):
    """Adds text where **bold** and *italic* spans are marked inline."""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**"):
            set_font(p.add_run(part[2:-2]), size, True, italic, color)
        elif part.startswith("*"):
            set_font(p.add_run(part[1:-1]), size, None, True, color)
        else:
            set_font(p.add_run(part), size, None, italic, color)


def fmt(p, before=0, after=6, line=1.3, align=None, keep=False, indent=None, first=None):
    f = p.paragraph_format
    f.space_before, f.space_after, f.line_spacing = Pt(before), Pt(after), line
    if align is not None:
        p.alignment = align
    if keep:
        f.keep_with_next = True
    if indent is not None:
        f.left_indent = indent
    if first is not None:
        f.first_line_indent = first
    return p


def field(p, instr, cached="", size=None, bold=None, color=None):
    """Inserts a complex field (PAGE, SEQ ...) with a cached result."""
    def r(tag, **attrs):
        run = p.add_run()
        set_font(run, size, bold, None, color)
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        if tag == "w:instrText":
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        run._element.append(el)
        return run
    r("w:fldChar", **{"w:fldCharType": "begin"})
    r("w:instrText")
    r("w:fldChar", **{"w:fldCharType": "separate"})
    set_font(p.add_run(cached), size, bold, None, color)
    r("w:fldChar", **{"w:fldCharType": "end"})


def fld_char(p, kind):
    run = p.add_run()
    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), kind)
    run._element.append(el)


def instr(p, text):
    run = p.add_run()
    el = OxmlElement("w:instrText")
    el.set(qn("xml:space"), "preserve")
    el.text = text
    run._element.append(el)


def shade(cell, hex_color):
    tcpr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hex_color)
    tcpr.append(s)


def cell_borders(cell, **sides):
    tcpr = cell._tc.get_or_add_tcPr()
    box = OxmlElement("w:tcBorders")
    for side, (style, sz, color) in sides.items():
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), style); b.set(qn("w:sz"), str(sz)); b.set(qn("w:space"), "0"); b.set(qn("w:color"), color)
        box.append(b)
    tcpr.append(box)


def table_borders(table, color="8C8C8C", size=4, inside=True):
    tblpr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        if side.startswith("inside") and not inside or color is None:
            e.set(qn("w:val"), "nil")
        else:
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(size)); e.set(qn("w:color"), color)
        b.append(e)
    tblpr.append(b)


def col_widths(table, widths_mm):
    table.autofit = False
    tblpr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblpr.append(lay)
    tw = tblpr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW"); tblpr.append(tw)
    tw.set(qn("w:type"), "dxa"); tw.set(qn("w:w"), str(int(sum(widths_mm) * 56.7)))
    for gc, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths_mm):
        gc.set(qn("w:w"), str(int(w * 56.7)))
    for row in table.rows:
        for c, w in zip(row.cells, widths_mm):
            c.width = Mm(w)


def cell_margins(table, top=50, bottom=50, left=90, right=90):
    tblpr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tblpr.append(m)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trpr.append(h)


def no_split(row):
    trpr = row._tr.get_or_add_trPr()
    c = OxmlElement("w:cantSplit"); c.set(qn("w:val"), "true"); trpr.append(c)


def bottom_rule(p, color=NAVY_HEX, size=8, space=2):
    ppr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(size)); e.set(qn("w:space"), str(space)); e.set(qn("w:color"), color)
    b.append(e); ppr.append(b)


def editable(p, text, tag, size=11.5, bold=False, italic=False, color=TEXT, align=None):
    """A Word content control (click and type) holding the given starter text."""
    sdt = OxmlElement("w:sdt")
    pr = OxmlElement("w:sdtPr")
    alias = OxmlElement("w:alias"); alias.set(qn("w:val"), tag)
    tg = OxmlElement("w:tag"); tg.set(qn("w:val"), tag.replace(" ", "_"))
    txt = OxmlElement("w:text")
    pr.extend([alias, tg, txt])
    content = OxmlElement("w:sdtContent")
    tmp = p.add_run(text)
    set_font(tmp, size, bold, italic, color)
    content.append(tmp._element)  # moves the run inside the control
    sdt.extend([pr, content])
    p._p.append(sdt)


# ---------------------------------------------------------------- document setup
def setup(doc):
    st = doc.styles
    normal = st["Normal"]
    normal.font.name = FONT; normal.font.size = Pt(12); normal.font.color.rgb = TEXT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for name, size, before, after in (("Heading 1", 16, 16, 8), ("Heading 2", 13, 10, 4), ("Heading 3", 12, 8, 4)):
        s = st[name]
        s.font.name = FONT; s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = NAVY
        s.font.italic = name == "Heading 3"
        s.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        rf = s.element.rPr.rFonts
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            rf.set(qn(a), FONT)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]
        pf = s.paragraph_format
        pf.space_before, pf.space_after, pf.keep_with_next, pf.line_spacing = Pt(before), Pt(after), True, 1.15

    cap = st["Caption"]
    cap.font.name = FONT; cap.font.size = Pt(10.5); cap.font.bold = True; cap.font.italic = False
    cap.font.color.rgb = NAVY

    for name, indent in (("toc 1", 0), ("toc 2", 6), ("table of figures", 0)):
        try:
            s = st[name]
        except KeyError:
            s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = normal
        s.font.name = FONT; s.font.size = Pt(11.5)
        s.font.bold = name == "toc 1"
        pf = s.paragraph_format
        pf.left_indent = Mm(indent); pf.space_before = Pt(2 if name == "toc 1" else 0); pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        pf.tab_stops.add_tab_stop(TEXT_W, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    # Ask Word to refresh the contents fields when the file is opened.
    settings = doc.settings.element
    uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true")
    settings.append(uf)


def page_setup(section, top=25, bottom=22):
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(25)
    section.top_margin, section.bottom_margin = Mm(top), Mm(bottom)
    section.header_distance = section.footer_distance = Mm(11)


def page_numbers(section, fmt_, start):
    pg = section._sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType"); section._sectPr.append(pg)
    pg.set(qn("w:fmt"), fmt_); pg.set(qn("w:start"), str(start))


def footer_number(section, with_header):
    section.footer.is_linked_to_previous = False
    section.header.is_linked_to_previous = False
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(fp, "PAGE", "1", 10.5, None, GREY)
    hp = section.header.paragraphs[0]
    if with_header:
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_font(hp.add_run("GDP Growth and Economic Well-Being"), 9.5, None, True, GREY)
        bottom_rule(hp, "A6A6A6", 4, 3)


# ---------------------------------------------------------------- builder
class Builder:
    def __init__(self, pages):
        self.pages = pages or {}
        self.doc = Document()
        setup(self.doc)
        self.fig_no = 0
        self.tab_no = 0
        self.box_no = 0
        self.toc = []      # (level, text)
        self.figs = []     # caption text
        self.tabs = []

    # ---- text blocks
    def h(self, level, text, new_page=False):
        if new_page:
            self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            self.doc.paragraphs[-1].paragraph_format.space_after = Pt(0)
        p = self.doc.add_heading(level=level)
        set_font(p.add_run(text), None, True, level == 3, NAVY)
        if level == 1:
            bottom_rule(p, NAVY_HEX, 8, 4)
        if level <= 2:
            self.toc.append((level, text))
        return p

    def p(self, text, after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        par = self.doc.add_paragraph()
        fmt(par, 0, after, 1.15, align)
        add_runs(par, text)
        return par

    def bullets(self, items, numbered=False):
        for i, it in enumerate(items, 1):
            par = self.doc.add_paragraph()
            fmt(par, 0, 2, 1.15, WD_ALIGN_PARAGRAPH.JUSTIFY, indent=Mm(9), first=Mm(-6))
            mark = f"{i}." if numbered else "▪"
            set_font(par.add_run(mark + "\t"), 12, numbered, None, NAVY)
            par.paragraph_format.tab_stops.add_tab_stop(Mm(9))
            add_runs(par, it)
        self.doc.paragraphs[-1].paragraph_format.space_after = Pt(8)

    def equation(self, text):
        par = self.doc.add_paragraph()
        fmt(par, 2, 8, 1.2, WD_ALIGN_PARAGRAPH.CENTER)
        set_font(par.add_run(text), 12, None, True, TEXT)

    def quote(self, text, who):
        par = self.doc.add_paragraph()
        fmt(par, 2, 2, 1.25, WD_ALIGN_PARAGRAPH.JUSTIFY, indent=Mm(12))
        par.paragraph_format.right_indent = Mm(12)
        set_font(par.add_run("“" + text + "”"), 11.5, None, True, TEXT)
        par2 = self.doc.add_paragraph()
        fmt(par2, 0, 8, 1.1, WD_ALIGN_PARAGRAPH.RIGHT)
        par2.paragraph_format.right_indent = Mm(12)
        set_font(par2.add_run("— " + who), 10.5, None, None, GREY)

    # ---- captions
    def caption(self, kind, title, before=4, after=2):
        par = self.doc.add_paragraph(style="Caption")
        fmt(par, before, after, 1.15, WD_ALIGN_PARAGRAPH.CENTER, keep=kind == "Table")
        set_font(par.add_run(f"{kind} "), 10.5, True, None, NAVY)
        n = self.fig_no if kind == "Figure" else self.tab_no
        field(par, f" SEQ {kind} \\* ARABIC ", str(n), 10.5, True, NAVY)
        set_font(par.add_run(f": {title}"), 10.5, True, None, NAVY)
        (self.figs if kind == "Figure" else self.tabs).append(f"{kind} {n}: {title}")

    def source(self, text):
        par = self.doc.add_paragraph()
        fmt(par, 0, 12, 1.1, WD_ALIGN_PARAGRAPH.CENTER)
        set_font(par.add_run("Source: "), 9.5, True, True, GREY)
        set_font(par.add_run(text), 9.5, None, True, GREY)

    def figure(self, name, title, src, width_mm=150):
        self.fig_no += 1
        par = self.doc.add_paragraph()
        fmt(par, 6, 0, 1.0, WD_ALIGN_PARAGRAPH.CENTER, keep=True)
        par.add_run().add_picture(str(FIG / f"{name}.png"), width=Mm(width_mm * 0.84))
        self.caption("Figure", title)
        self.source(src)

    def table(self, title, headers, rows, widths, src=None, size=10, align_first_left=True, bold_first=False):
        self.tab_no += 1
        self.caption("Table", title, before=8, after=4)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_borders(t, "A6A6A6", 4)
        col_widths(t, widths)
        cell_margins(t)
        for j, htxt in enumerate(headers):
            c = t.cell(0, j); shade(c, NAVY_HEX)
            cp = c.paragraphs[0]; fmt(cp, 1, 1, 1.1, WD_ALIGN_PARAGRAPH.CENTER)
            set_font(cp.add_run(htxt), size, True, None, RGBColor(0xFF, 0xFF, 0xFF))
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        repeat_header(t.rows[0])
        for i, row in enumerate(rows, 1):
            no_split(t.rows[i])
            for j, val in enumerate(row):
                c = t.cell(i, j)
                if i % 2 == 0:
                    shade(c, "F2F5FA")
                cp = c.paragraphs[0]
                left = j == 0 and align_first_left or len(str(val)) > 22
                fmt(cp, 1, 1, 1.1, WD_ALIGN_PARAGRAPH.LEFT if left else WD_ALIGN_PARAGRAPH.CENTER)
                add_runs(cp, str(val), size)
                if j == 0 and bold_first:
                    for r in cp.runs:
                        r.bold = True
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if src:
            self.source(src)
        else:
            sp = self.doc.add_paragraph(); fmt(sp, 0, 8, 1.0)

    def demo(self, title, lines):
        """Shaded worked-example box."""
        self.box_no += 1
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_borders(t, None)
        col_widths(t, [160])
        cell_margins(t, 120, 120, 200, 200)
        c = t.cell(0, 0); shade(c, PALE_HEX)
        cell_borders(c, left=("single", 24, NAVY_HEX))
        no_split(t.rows[0])
        cp = c.paragraphs[0]; fmt(cp, 0, 4, 1.15, keep=True)
        set_font(cp.add_run(f"Demonstration {self.box_no}: {title}"), 11.5, True, None, NAVY)
        for ln in lines:
            par = c.add_paragraph()
            if ln.startswith("= "):
                fmt(par, 1, 1, 1.15, WD_ALIGN_PARAGRAPH.CENTER)
                set_font(par.add_run(ln[2:]), 11, None, True, TEXT)
            else:
                fmt(par, 0, 3, 1.2, WD_ALIGN_PARAGRAPH.JUSTIFY)
                add_runs(par, ln, 11)
        sp = self.doc.add_paragraph(); fmt(sp, 0, 6, 1.0)

    # ---- contents lists (real Word fields with cached entries)
    def contents_field(self, instr_text, entries, style_for):
        first = True
        for i, (level, text) in enumerate(entries):
            par = self.doc.add_paragraph(style=style_for(level))
            if first:
                fld_char(par, "begin"); instr(par, instr_text); fld_char(par, "separate")
                first = False
            pg = self.pages.get(text, "")
            set_font(par.add_run(text), 11, level == 1, None, TEXT)
            set_font(par.add_run("\t" + str(pg)), 11, level == 1, None, TEXT)
            if i == len(entries) - 1:
                fld_char(par, "end")

    def front_heading(self, text):
        par = self.doc.add_paragraph()
        fmt(par, 0, 14, 1.0, WD_ALIGN_PARAGRAPH.CENTER)
        set_font(par.add_run(text), 16, True, None, NAVY, caps=True)
        bottom_rule(par, NAVY_HEX, 8, 6)
        return par


# ---------------------------------------------------------------- cover page
def cover(b):
    doc = b.doc
    sec = doc.sections[0]
    page_setup(sec, top=14, bottom=12)
    sec.left_margin = sec.right_margin = Mm(20)
    pg = OxmlElement("w:pgBorders"); pg.set(qn("w:offsetFrom"), "page")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "thinThickSmallGap"); e.set(qn("w:sz"), "18"); e.set(qn("w:space"), "22"); e.set(qn("w:color"), NAVY_HEX)
        pg.append(e)
    sec._sectPr.append(pg)

    p = doc.add_paragraph(); fmt(p, 0, 2, 1.0, WD_ALIGN_PARAGRAPH.CENTER)
    p.add_run().add_picture(str(HERE / "diu-seal.png"), width=Mm(20))
    p = doc.add_paragraph(); fmt(p, 0, 0, 1.0, WD_ALIGN_PARAGRAPH.CENTER)
    editable(p, "Dhaka International University", "University Name", 17, True, color=NAVY)
    p = doc.add_paragraph(); fmt(p, 2, 6, 1.0, WD_ALIGN_PARAGRAPH.CENTER)
    editable(p, "Department of Business Administration", "Department", 11.5, False, True, GREY)

    p = doc.add_paragraph(); fmt(p, 0, 6, 1.0, WD_ALIGN_PARAGRAPH.CENTER)
    p.add_run().add_picture(str(FIG / "cover.png"), width=Mm(158))

    # Course line
    t = doc.add_table(rows=2, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t, None); col_widths(t, [40, 118]); cell_margins(t, 30, 30, 80, 80)
    for i, (k, v) in enumerate((("Course Title", "[Write course title, e.g. Macroeconomics]"),
                                ("Course Code", "[Write course code, e.g. ECO-201]"))):
        kp = t.cell(i, 0).paragraphs[0]; fmt(kp, 2, 2, 1.0)
        set_font(kp.add_run(k + " :"), 11, True, None, NAVY)
        vp = t.cell(i, 1).paragraphs[0]; fmt(vp, 2, 2, 1.0)
        editable(vp, v, k, 11.5)
        cell_borders(t.cell(i, 1), bottom=("dotted", 4, "8C8C8C"))

    sp = doc.add_paragraph(); fmt(sp, 0, 6, 1.0)

    to_rows = [("Name", "[Teacher's Name]"), ("Designation", "[Lecturer / Assistant Professor]"),
               ("Department", "[Department Name]"), ("University", "[University Name]")]
    by_rows = [("Name", "[Your Full Name]"), ("Student ID", "[Your ID]"), ("Batch", "[Batch]"),
               ("Section", "[Section]"), ("Semester", "[Semester]"), ("Program", "[BBA / BSS / BSc]")]
    n = max(len(to_rows), len(by_rows))
    t = doc.add_table(rows=n + 1, cols=5)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t, None); col_widths(t, [22, 52, 8, 24, 54]); cell_margins(t, 20, 20, 60, 60)
    for col, title in ((0, "Submitted To"), (3, "Submitted By")):
        a = t.cell(0, col).merge(t.cell(0, col + 1)); shade(a, NAVY_HEX)
        ap = a.paragraphs[0]; fmt(ap, 2, 2, 1.0, WD_ALIGN_PARAGRAPH.CENTER)
        set_font(ap.add_run(title), 12, True, None, RGBColor(0xFF, 0xFF, 0xFF), caps=True)
    for rows, col in ((to_rows, 0), (by_rows, 3)):
        for i, (k, v) in enumerate(rows, 1):
            kp = t.cell(i, col).paragraphs[0]; fmt(kp, 3, 1, 1.0)
            set_font(kp.add_run(k), 10, True, None, NAVY)
            vp = t.cell(i, col + 1).paragraphs[0]; fmt(vp, 3, 1, 1.0)
            editable(vp, v, ("Teacher " if col == 0 else "Student ") + k, 11)
            cell_borders(t.cell(i, col + 1), bottom=("dotted", 4, "8C8C8C"))
            t.cell(i, col).vertical_alignment = t.cell(i, col + 1).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM

    p = doc.add_paragraph(); fmt(p, 12, 0, 1.0, WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p.add_run("Date of Submission :  "), 11.5, True, None, NAVY)
    editable(p, "[DD Month YYYY]", "Submission Date", 11.5)


# ---------------------------------------------------------------- whole document
def build(pages):
    b = Builder(pages)
    cover(b)
    doc = b.doc

    # front matter section (roman numbers)
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    page_setup(sec)
    pgb = sec._sectPr.find(qn("w:pgBorders"))
    if pgb is not None:
        sec._sectPr.remove(pgb)
    page_numbers(sec, "lowerRoman", 1)
    footer_number(sec, False)

    b.front_heading("Abstract")
    for para in ABSTRACT:
        b.p(para, after=8)
    kp = doc.add_paragraph(); fmt(kp, 6, 0, 1.2, WD_ALIGN_PARAGRAPH.JUSTIFY)
    set_font(kp.add_run("Keywords: "), 11.5, True, None, NAVY)
    set_font(kp.add_run("GDP, economic growth, economic well-being, Human Development Index, Genuine Progress "
                        "Indicator, inequality, sustainability, beyond GDP, Bangladesh"), 11.5, None, True)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    b.front_heading("Table of Contents")
    toc_entries = pages.get("__toc__", [(1, "1. Introduction")])
    b.contents_field('TOC \\o "1-2" \\z \\u', toc_entries, lambda lv: "toc 1" if lv == 1 else "toc 2")

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    b.front_heading("List of Figures")
    b.contents_field('TOC \\h \\z \\c "Figure"', [(0, t) for t in pages.get("__figs__", ["Figure 1"])],
                     lambda lv: "table of figures")
    sp = doc.add_paragraph(); fmt(sp, 0, 10, 1.0)
    b.front_heading("List of Tables")
    b.contents_field('TOC \\h \\z \\c "Table"', [(0, t) for t in pages.get("__tabs__", ["Table 1"])],
                     lambda lv: "table of figures")

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    b.front_heading("List of Abbreviations")
    t = doc.add_table(rows=len(ABBREVIATIONS), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t, None); col_widths(t, [30, 120]); cell_margins(t, 25, 25, 80, 80)
    for i, (a, d) in enumerate(ABBREVIATIONS):
        ap = t.cell(i, 0).paragraphs[0]; fmt(ap, 1, 1, 1.15)
        set_font(ap.add_run(a), 11.5, True, None, NAVY)
        dp = t.cell(i, 1).paragraphs[0]; fmt(dp, 1, 1, 1.15)
        set_font(dp.add_run(d), 11.5)
        cell_borders(t.cell(i, 0), bottom=("dotted", 4, "BFBFBF")); cell_borders(t.cell(i, 1), bottom=("dotted", 4, "BFBFBF"))

    # main body section (arabic numbers)
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    page_setup(sec)
    page_numbers(sec, "decimal", 1)
    footer_number(sec, True)

    first_chapter = True
    for block in BODY:
        kind = block[0]
        if kind == "h1":
            b.h(1, block[1], new_page=False)
            first_chapter = False
        elif kind == "h2":
            b.h(2, block[1])
        elif kind == "h3":
            b.h(3, block[1])
        elif kind == "p":
            b.p(block[1])
        elif kind == "ul":
            b.bullets(block[1])
        elif kind == "ol":
            b.bullets(block[1], numbered=True)
        elif kind == "eq":
            b.equation(block[1])
        elif kind == "quote":
            b.quote(block[1], block[2])
        elif kind == "fig":
            b.figure(*block[1:])
        elif kind == "table":
            b.table(**block[1])
        elif kind == "demo":
            b.demo(block[1], block[2])
        elif kind == "page":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # references
    b.h(1, "References", new_page=True)
    for ref in REFERENCES:
        par = doc.add_paragraph()
        fmt(par, 0, 6, 1.2, WD_ALIGN_PARAGRAPH.LEFT, indent=Mm(10), first=Mm(-10))
        add_runs(par, ref, 11)

    b.h(1, "Appendix A: Glossary of Key Terms", new_page=True)
    b.table("Glossary of key terms used in the assignment", ["Term", "Meaning"], GLOSSARY, [42, 118],
            None, 10.5, bold_first=True)
    return b


def pdf_pages(docx_path):
    tmp = Path(tempfile.mkdtemp())
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(docx_path)],
                   check=True, capture_output=True)
    pdf = tmp / (docx_path.stem + ".pdf")
    txt = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], check=True, capture_output=True, text=True).stdout
    return pdf, [re.sub(r"\s+", " ", pg) for pg in txt.split("\f")]


def norm(s):
    return re.sub(r"\s+", " ", s).strip()


def locate(b, pages_text):
    """Finds the printed (arabic) page of every heading and caption from a rendered PDF."""
    body_start = next(i for i, t in enumerate(pages_text)
                      if "1. Introduction" in t and "table of contents" not in t.lower() and "list of" not in t.lower())
    found = {}
    items = [("h", t) for _, t in b.toc]
    # Captions appear in document order interleaved with headings; search them separately.
    def scan(seq):
        cur = body_start
        for text in seq:
            key = norm(text)[:48]
            for i in range(cur, len(pages_text)):
                if key in pages_text[i]:
                    found[text] = i - body_start + 1
                    cur = i
                    break
    scan([t for _, t in items])
    scan(b.figs)
    scan(b.tabs)
    return found


def main():
    b = build({})
    b.doc.save(OUT)
    _, txt = pdf_pages(OUT)
    pages = locate(b, txt)
    pages["__toc__"] = b.toc
    pages["__figs__"] = b.figs
    pages["__tabs__"] = b.tabs
    for _ in range(2):  # second pass can shift the front matter by a page, so settle twice
        b2 = build(pages)
        b2.doc.save(OUT)
        pdf, txt = pdf_pages(OUT)
        new = locate(b2, txt)
        new.update({k: pages[k] for k in ("__toc__", "__figs__", "__tabs__")})
        if new == pages:
            break
        pages = new
    missing = [t for _, t in b.toc if t not in pages] + [t for t in b.figs + b.tabs if t not in pages]
    print(f"saved {OUT.name}: {len(txt) - 1} pages, {b.fig_no} figures, {b.tab_no} tables")
    if missing:
        print("not located:", missing, file=sys.stderr)
    subprocess.run(["cp", str(pdf), str(OUT.with_suffix(".pdf"))], check=True)


if __name__ == "__main__":
    main()
