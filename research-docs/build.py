"""Build the technical document as PDF (Chromium) and DOCX (pandoc) from src/*.md.

Usage:  python3 research-docs/build.py
Needs:  pandoc >= 3, node with the `playwright` package, and MathJax 3 installed locally:
        (cd research-docs/build && npm install mathjax@3.2.2)
"""
import os
import re
import shutil
import subprocess
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Cm

HERE = Path(__file__).resolve().parent
SRC, BUILD, FIG = HERE / "src", HERE / "build", HERE / "figures"
NAME = "Defect_Aware_MultiDie_Fusion_Technical_Document"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)


def run(cmd, **kw):
    subprocess.run(cmd, check=True, **kw)


def combined_markdown():
    parts = [p.read_text(encoding="utf-8") for p in sorted(SRC.glob("*.md"))]
    md = "\n\n".join(parts)
    (BUILD / "document.md").write_text(md, encoding="utf-8")
    return md


def build_html_pdf():
    mathjax = BUILD / "node_modules/mathjax/es5/tex-svg-full.js"
    if not mathjax.exists():
        raise SystemExit("MathJax missing: run (cd research-docs/build && npm install mathjax@3.2.2)")
    html = BUILD / f"{NAME}.html"
    run(["pandoc", str(BUILD / "document.md"), "-s", "-o", str(html), "--toc", "--toc-depth=2",
         f"--mathjax={mathjax.as_uri()}", f"--css={(HERE / 'style.css').as_uri()}",
         "--metadata", "pagetitle=Multi-Die Reference Reconstruction",
         f"--resource-path={HERE}"])
    # Make figure paths absolute so Chromium resolves them from build/
    text = html.read_text(encoding="utf-8").replace('src="figures/', f'src="{FIG.as_uri()}/')
    # pandoc's template adds a remote polyfill script for MathJax; the build must stay offline
    text = re.sub(r'<script src="https://polyfill\.io[^"]*"></script>\s*', "", text)
    # MathJax config: SVG output with fonts as paths, no menu
    cfg = ("<script>window.MathJax={svg:{fontCache:'global'},options:{enableMenu:false},"
           "startup:{typeset:true}};</script>")
    text = text.replace("<script src=", cfg + "<script src=", 1)
    html.write_text(text, encoding="utf-8")

    js = BUILD / "print.cjs"
    js.write_text(r"""
const { chromium } = require('playwright');
const { pathToFileURL } = require('node:url');
(async () => {
  const [src, out] = process.argv.slice(2);
  const b = await chromium.launch();
  const p = await b.newPage();
  await p.goto(pathToFileURL(src).href, { waitUntil: 'networkidle' });
  await p.waitForFunction(() => window.MathJax && MathJax.startup && MathJax.startup.promise);
  await p.evaluate(() => MathJax.startup.promise);
  const errs = await p.evaluate(() => document.querySelectorAll('mjx-merror').length);
  if (errs) { console.error(`MathJax errors: ${errs}`); process.exit(1); }
  await p.emulateMedia({ media: 'print' });
  await p.pdf({ path: out, format: 'A4', printBackground: true, preferCSSPageSize: true, tagged: true, outline: true,
    displayHeaderFooter: true,
    headerTemplate: '<div style="font-family:Liberation Sans,Arial;font-size:7.5pt;color:#5b6573;width:100%;padding:0 15mm;">Defect-Aware Multi-Die Image Fusion · Research and implementation documentation</div>',
    footerTemplate: '<div style="font-family:Liberation Sans,Arial;font-size:7.5pt;color:#5b6573;width:100%;padding:0 15mm;text-align:right;">Page <span class="pageNumber"></span> of <span class="totalPages"></span></div>',
    margin: { top: '16mm', bottom: '15mm', left: '15mm', right: '15mm' } });
  await b.close();
})();
""")
    npm_root = subprocess.run(["npm", "root", "-g"], capture_output=True, text=True).stdout.strip()
    run(["node", str(js), str(html), str(HERE / f"{NAME}.pdf")], env={**os.environ, "NODE_PATH": npm_root})


def style(doc, name):
    """Find a style by its displayed name (pandoc stores 'Heading 1', which python-docx's
    lookup translates to the built-in id and misses)."""
    for s in doc.styles:
        if s.name == name:
            return s
    return None


def reference_docx():
    ref = BUILD / "reference.docx"
    run(["pandoc", "-o", str(ref), "--print-default-data-file", "reference.docx"])
    d = Document(ref)
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if s := style(d, name):
            s.font.name = "Liberation Serif"
            s.font.size = Pt(10.5)
    for name, size in (("Title", 24), ("Subtitle", 14), ("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11)):
        if s := style(d, name):
            s.font.name = "Liberation Sans"
            s.font.size = Pt(size)
            s.font.color.rgb = ACCENT
            s.font.bold = name != "Subtitle"
    for name in ("Author", "Date"):
        if s := style(d, name):
            s.font.name = "Liberation Sans"
            s.font.size = Pt(11)
    for name in ("Image Caption", "Table Caption"):
        if s := style(d, name):
            s.font.name = "Liberation Sans"
            s.font.size = Pt(8.5)
            s.font.italic = False
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.9))
    sec.top_margin = sec.bottom_margin = Cm(1.9)
    d.save(ref)
    return ref


# Child-element order required by the WordprocessingML schema (ECMA-376), for the
# containers where pandoc's output deviates from it.
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
SEQ = {
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs "
           "suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd "
           "snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment "
           "textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid "
           "vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em "
           "lang eastAsianLayout specVanish oMath rPrChange",
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd "
             "tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription tblPrChange",
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark headers "
            "cellIns cellDel cellMerge tcPrChange",
    "style": "name aliases basedOn next link autoRedefine hidden uiPriority semiHidden unhideWhenUsed qFormat locked personal "
             "personalCompose personalReply rsid pPr rPr tblPr trPr tcPr tblStylePr",
    "tblStylePr": "pPr rPr tblPr trPr tcPr",
    "sectPr": "headerReference footerReference footnotePr endnotePr type pgSz pgMar paperSrc pgBorders lnNumType pgNumType "
              "cols formProt vAlign noEndnote titlePg textDirection bidi rtlGutter docGrid printerSettings sectPrChange",
    "settings": "writeProtection view zoom removePersonalInformation removeDateAndTime doNotDisplayPageBoundaries "
                "displayBackgroundShape printPostScriptOverText printFractionalCharacterWidth printFormsData "
                "embedTrueTypeFonts embedSystemFonts saveSubsetFonts saveFormsData mirrorMargins alignBordersAndEdges "
                "bordersDoNotSurroundHeader bordersDoNotSurroundFooter gutterAtTop hideSpellingErrors "
                "hideGrammaticalErrors activeWritingStyle proofState formsDesign attachedTemplate linkStyles "
                "stylePaneFormatFilter stylePaneSortMethod documentType mailMerge revisionView trackRevisions "
                "doNotTrackMoves doNotTrackFormatting documentProtection autoFormatOverride styleLockTheme "
                "styleLockQFSet defaultTabStop autoHyphenation consecutiveHyphenLimit hyphenationZone doNotHyphenateCaps "
                "showEnvelope summaryLength clickAndTypeStyle defaultTableStyle evenAndOddHeaders bookFoldRevPrinting "
                "bookFoldPrinting bookFoldPrintingSheets drawingGridHorizontalSpacing drawingGridVerticalSpacing "
                "displayHorizontalDrawingGridEvery displayVerticalDrawingGridEvery doNotUseMarginsForDrawingGridOrigin "
                "drawingGridHorizontalOrigin drawingGridVerticalOrigin doNotShadeFormData noPunctuationKerning "
                "characterSpacingControl printTwoOnOne strictFirstAndLastChars noLineBreaksAfter noLineBreaksBefore "
                "savePreviewPicture doNotValidateAgainstSchema saveInvalidXml ignoreMixedContent "
                "alwaysShowPlaceholderText doNotDemarcateInvalidXml saveXmlDataOnly useXSLTWhenSaving saveThroughXslt "
                "showXMLTags alwaysMergeEmptyNamespace updateFields hdrShapeDefaults footnotePr endnotePr compat docVars "
                "rsids mathPr attachedSchema themeFontLang clrSchemeMapping doNotIncludeSubdocsInStats "
                "doNotAutoCompressPictures forceUpgrade captions readModeInkLockDown smartTagType schemaLibrary "
                "shapeDefaults doNotEmbedSmartTags decimalSymbol listSeparator",
}
M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
SEQ_M = {"mcPr": "count mcJc", "rPr": "lit nor scr sty brk aln",   # math (m:) containers
         "dPr": "begChr sepChr endChr grow shp ctrlPr"}
ORDER = {k: {n: i for i, n in enumerate(v.split())} for k, v in SEQ.items()}
ORDER_M = {k: {n: i for i, n in enumerate(v.split())} for k, v in SEQ_M.items()}


def local_name(tag):
    return tag.split("}", 1)[1] if isinstance(tag, str) and "}" in tag else ""


def reorder(el, rank):
    kids = list(el)
    ordered = sorted(kids, key=lambda c: rank.get(local_name(c.tag), len(rank)))
    if ordered != kids:
        for c in kids:
            el.remove(c)
        el.extend(ordered)


def schema_fix(xml_bytes, part):
    """Reorder children to schema order and fill pandoc's known omissions in one XML part."""
    from lxml import etree
    root = etree.fromstring(xml_bytes)
    for el in list(root.iter()):
        if not isinstance(el.tag, str):
            continue
        local = local_name(el.tag)
        if el.tag.startswith(M):
            if local == "rPr" and el.find(M + "nor") is not None:
                for c in el.findall(M + "sty") + el.findall(M + "scr"):   # nor excludes scr/sty
                    el.remove(c)
            if local in ORDER_M:
                reorder(el, ORDER_M[local])
            continue
        if not el.tag.startswith(W):
            continue
        if local == "rPr":                       # pandoc leaves stray text inside some w:rPr
            if el.text and el.text.strip():
                el.text = None
            for c in el:
                if c.tail and c.tail.strip():
                    c.tail = None
        if local in ORDER:
            reorder(el, ORDER[local])
        if local == "pgMar":
            for attr, default in (("header", "720"), ("footer", "720"), ("gutter", "0")):
                if el.get(W + attr) is None:
                    el.set(W + attr, default)
        if local == "nsid":
            v = el.get(W + "val", "")
            if len(v) < 8:
                el.set(W + "val", v.rjust(8, "0"))
    if part == "word/settings.xml" and root.find(W + "updateFields") is None:
        uf = etree.SubElement(root, W + "updateFields")
        uf.set(W + "val", "true")
        reorder(root, ORDER["settings"])
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def rebalance_pipe_tables(md):
    """Pandoc sizes DOCX columns from the dashes in a pipe table's separator row.
    Rewrite each separator so column widths follow the content (sqrt-damped mean length)."""
    lines, out, i = md.split("\n"), [], 0
    sep = re.compile(r"^\|(\s*:?-+:?\s*\|)+\s*$")
    while i < len(lines):
        if i + 1 < len(lines) and lines[i].startswith("|") and sep.match(lines[i + 1]):
            j = i + 2
            while j < len(lines) and lines[j].startswith("|"):
                j += 1
            rows = [[c.strip() for c in l.strip().strip("|").split("|")] for l in [lines[i]] + lines[i + 2:j]]
            ncol = len(rows[0])
            lens = []
            for c in range(ncol):
                vals = [len(re.sub(r"<[^>]+>|\*|`", "", r[c])) for r in rows if c < len(r)]
                lens.append(max(4.0, (sum(vals) / len(vals)) ** 0.5 * 1.0 + 0.15 * max(vals) ** 0.5))
            total = sum(lens)
            dashes = [max(3, round(72 * l / total)) for l in lens]
            out += [lines[i], "|" + "|".join("-" * d for d in dashes) + "|"] + lines[i + 2:j]
            i = j
            continue
        out.append(lines[i])
        i += 1
    return "\n".join(out)


def build_docx():
    md = (BUILD / "document.md").read_text(encoding="utf-8")
    md = re.sub(r"\(figures/(fig\w+)\.svg\)", r"(figures/\1.png)", md)
    md = rebalance_pipe_tables(md)
    docx_md = BUILD / "document_docx.md"
    docx_md.write_text(md, encoding="utf-8")
    out = HERE / f"{NAME}.docx"
    run(["pandoc", str(docx_md), "-o", str(out), "--toc", "--toc-depth=2", "--columns=72",
         f"--reference-doc={reference_docx()}", f"--resource-path={HERE}"])
    post_process_docx(out)


def post_process_docx(path):
    """Style tables, fix schema order, and ask Word to refresh the table of contents on open."""
    d = Document(path)
    for t in d.tables:
        t.style = style(d, "Table") or t.style
        for i, row in enumerate(t.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(8)
                        r.font.name = "Liberation Sans"
                        if i == 0:
                            r.font.bold = True
                            r.font.color.rgb = ACCENT
    for p in d.paragraphs:
        if p.style.name == "Image Caption":
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    d.save(path)
    tmp = path.with_suffix(".tmp.docx")
    parts = {"word/document.xml", "word/styles.xml", "word/numbering.xml", "word/settings.xml"}
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in parts:
                data = schema_fix(data, item.filename)
            zout.writestr(item, data)
    shutil.move(tmp, path)


def build_summary_pdf():
    """Print decisions_summary.html to a one-page PDF; fail if it overflows."""
    js = BUILD / "summary.cjs"
    js.write_text(r"""
const { chromium } = require('playwright');
const { pathToFileURL } = require('node:url');
(async () => {
  const [src, out] = process.argv.slice(2);
  const b = await chromium.launch();
  const p = await b.newPage({ viewport: { width: 794, height: 1123 } });
  await p.goto(pathToFileURL(src).href, { waitUntil: 'networkidle' });
  await p.emulateMedia({ media: 'print' });
  const over = await p.evaluate(() => { const g = document.querySelector('.page'); return g.scrollHeight - g.clientHeight; });
  if (over > 1) { console.error(`summary overflows by ${over}px`); process.exit(1); }
  await p.pdf({ path: out, format: 'A4', printBackground: true, preferCSSPageSize: true, tagged: true,
                margin: { top: '0', right: '0', bottom: '0', left: '0' } });
  await b.close();
})();
""")
    npm_root = subprocess.run(["npm", "root", "-g"], capture_output=True, text=True).stdout.strip()
    run(["node", str(js), str(HERE / "decisions_summary.html"), str(HERE / "Decisions_and_Open_Questions.pdf")],
        env={**os.environ, "NODE_PATH": npm_root})


if __name__ == "__main__":
    BUILD.mkdir(exist_ok=True)
    combined_markdown()
    build_html_pdf()
    build_docx()
    build_summary_pdf()
    print(f"Wrote {NAME}.pdf, {NAME}.docx and Decisions_and_Open_Questions.pdf")
