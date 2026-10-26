# Builds the editable Word copy of the DIU BBA cover page (same layout as the HTML version).
# Usage: python3 cover-page/build_docx.py   (needs python-docx)
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

HERE = Path(__file__).parent
OUT = HERE / "DIU_BBA_Cover_Page.docx"

INK = RGBColor(0x0F, 0x4C, 0x3A)
BRASS = "A8843A"
INK_HEX = "0F4C3A"
TEXT = RGBColor(0x1C, 0x24, 0x21)
MUTED = RGBColor(0x5D, 0x6A, 0x64)
HINT = RGBColor(0x8E, 0x9A, 0x95)
SERIF = "Georgia"
SANS = "Calibri"
BANGLA = "Nirmala UI"


def run(p, text, size, font=SERIF, color=TEXT, bold=False, italic=False, spacing=0, caps=False):
    r = p.add_run(text)
    r.font.name = font
    r._element.rPr.rFonts.set(qn("w:cs"), font)
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.bold, r.italic, r.font.all_caps = bold, italic, caps
    if spacing:
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(int(spacing * 20)))
        r._element.rPr.append(sp)
    return r


def para(container, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=0):
    p = container.add_paragraph()
    p.alignment = align
    f = p.paragraph_format
    f.space_before, f.space_after, f.line_spacing = Pt(before), Pt(after), 1.0
    return p


def borders(el, tag, **sides):
    """sides: top/bottom/left/right = (style, size_eighths_pt, hex)"""
    pr = el.get_or_add_pPr() if tag == "w:pBdr" else el
    box = OxmlElement(tag)
    for side, (style, sz, color) in sides.items():
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), style)
        b.set(qn("w:sz"), str(sz))
        b.set(qn("w:space"), "4" if tag == "w:pBdr" else "0")
        b.set(qn("w:color"), color)
        box.append(b)
    pr.append(box)


def cell_borders(cell, **sides):
    borders(cell._tc.get_or_add_tcPr(), "w:tcBorders", **sides)


def shade(cell, hex_color):
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(s)


def no_table_borders(table):
    tbl_pr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "nil")
        b.append(e)
    tbl_pr.append(b)


def widths(table, *mm):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(sum(mm) * 56.7)))
    for col, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), mm):
        col.set(qn("w:w"), str(int(w * 56.7)))
    for row in table.rows:
        for cell, w in zip(row.cells, mm):
            cell.width = Mm(w)


def cell_text(cell, text, size, font=SERIF, color=TEXT, align=WD_ALIGN_PARAGRAPH.LEFT, **kw):
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    run(p, text, size, font, color, **kw)
    return p


def label(cell, text):
    cell_text(cell, text, 7.5, SANS, MUTED, bold=True, spacing=1, caps=True)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM


def value(cell, text, hint=False, big=False):
    cell_text(cell, text, 12 if big else 11, SERIF, HINT if hint else TEXT, italic=hint, bold=big and not hint)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    cell_borders(cell, bottom=("dotted", 4, "9AA5A0"))


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin = sec.bottom_margin = Mm(16)
sec.left_margin = sec.right_margin = Mm(22)

# Double page frame in the main colour.
pg = OxmlElement("w:pgBorders")
pg.set(qn("w:offsetFrom"), "page")
for side in ("top", "left", "bottom", "right"):
    b = OxmlElement(f"w:{side}")
    b.set(qn("w:val"), "thinThickSmallGap")
    b.set(qn("w:sz"), "18")
    b.set(qn("w:space"), "26")
    b.set(qn("w:color"), INK_HEX)
    pg.append(b)
sec._sectPr.append(pg)

# --- header ---
para(doc, after=4).add_run().add_picture(str(HERE / "diu-seal.png"), width=Mm(30))

run(para(doc, after=2), "Dhaka International University", 19, SERIF, INK, caps=True, spacing=0.4)
run(para(doc, after=8), "ঢাকা ইন্টারন্যাশনাল ইউনিভার্সিটি", 12.5, BANGLA, MUTED)
p = para(doc, after=3)
run(p, "Department of Business Administration", 9.5, SANS, TEXT, bold=True, caps=True, spacing=2)
borders(p._p, "w:pBdr", top=("single", 4, BRASS))
run(para(doc), "Faculty of Business Studies", 11, SERIF, MUTED, italic=True)

# --- the work ---
para(doc, before=40)
t = doc.add_table(rows=1, cols=1)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
no_table_borders(t)
widths(t, 70)
c = t.cell(0, 0)
shade(c, INK_HEX)
cp = cell_text(c, "Assignment", 15, SERIF, RGBColor(0xFF, 0xFF, 0xFF), align=WD_ALIGN_PARAGRAPH.CENTER, caps=True, spacing=6)
cp.paragraph_format.space_before = cp.paragraph_format.space_after = Pt(6)

run(para(doc, before=10, after=8), "on", 11, SERIF, MUTED, italic=True)
run(para(doc, after=14), "Write your assignment topic here", 21, SERIF, HINT, italic=True)

t = doc.add_table(rows=2, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
no_table_borders(t)
widths(t, 34, 94)
for i, (k, v) in enumerate([("Course Title", "e.g. Principles of Management"), ("Course Code", "e.g. MGT-101")]):
    cell_text(t.cell(i, 0), k, 8.5, SANS, INK, bold=True, caps=True, spacing=1)
    cell_text(t.cell(i, 1), v, 12, SERIF, HINT, italic=True)
for c in t.rows[0].cells:
    cell_borders(c, top=("single", 4, "1C2421"))
for c in t.rows[1].cells:
    cell_borders(c, bottom=("double", 6, "1C2421"))  # accountant's double rule

# --- submitted to / by ---
para(doc, before=48)
to_rows = [("Name", "Teacher's name", True, True), ("Designation", "Lecturer / Assistant Professor", True, False),
           ("Department", "Business Administration", False, False), ("Institution", "Dhaka International University", False, False)]
by_rows = [("Name", "Your full name", True, True), ("Roll / ID", "Your roll or student ID", True, False),
           ("Reg. No.", "Registration number", True, False), ("Batch", "e.g. 72nd", True, False),
           ("Section", "e.g. A", True, False), ("Semester", "e.g. Fall 2026", True, False), ("Program", "BBA", False, False)]
t = doc.add_table(rows=1 + len(by_rows), cols=5)
no_table_borders(t)
widths(t, 23, 50, 6, 23, 64)
for col, title in ((0, "Submitted To"), (3, "Submitted By")):
    a = t.cell(0, col).merge(t.cell(0, col + 1))
    p = cell_text(a, "◆  ", 9, SERIF, RGBColor.from_string(BRASS))
    run(p, title, 13, SERIF, INK, caps=True, spacing=1.5)
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(8), Pt(4)
for c in t.rows[0].cells:
    cell_borders(c, top=("single", 4, BRASS))
for i, (k, v, hint, big) in enumerate(to_rows, start=1):
    label(t.cell(i, 0), k)
    value(t.cell(i, 1), v, hint, big)
for i, (k, v, hint, big) in enumerate(by_rows, start=1):
    label(t.cell(i, 3), k)
    value(t.cell(i, 4), v, hint, big)
for row in t.rows:
    cell_borders(row.cells[2], right=("single", 4, BRASS))

# --- date + address ---
p = para(doc, before=14, after=0)
run(p, "Date of Submission:   ", 8.5, SANS, INK, bold=True, caps=True, spacing=1.5)
run(p, "DD / MM / YYYY", 12, SERIF, HINT, italic=True)
p.paragraph_format.space_before = Pt(16)
borders(p._p, "w:pBdr", top=("single", 4, "1C2421"), bottom=("double", 6, "1C2421"))
run(para(doc, before=8), "Satarkul, Badda, Dhaka-1212, Bangladesh", 7.8, SANS, MUTED, caps=True, spacing=1.2)

doc.save(OUT)
print("wrote", OUT)
